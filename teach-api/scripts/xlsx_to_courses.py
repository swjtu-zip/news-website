#!/usr/bin/env python3
"""把教务导出的选修课 Excel(“课程清单”工作表)转换成 teach-api 的课程快照 JSON。

用法:
    python3 xlsx_to_courses.py courses整理.xlsx courses.json

输出格式见 teach-api internal/store 的 CoursesSnapshot。教学班号是主键,
同一表格中必须唯一;教师姓名原样保留,与教师目录的关联由 teach-api 导入时完成。
"""

import json
import sys
from datetime import datetime, timezone

import openpyxl

COLUMNS = {
    "教学班号": "id",
    "学期": "term",
    "开课学院": "college",
    "课程代码": "course_code",
    "课程名称": "course_name",
    "教学班": "class_num",
    "教师": "teacher_name",
    "职称": "teacher_title",
    "学分": "credit",
    "总学时": "hours_total",
    "周学时": "hours_week",
    "周次": "weeks",
    "星期": "weekday",
    "节次": "periods",
    "校区": "campus",
    "上课时间": "schedule_text",
    "考核方式": "assessment",
    "课程性质": "nature",
    "课程分类": "category",
    "备注": "remark",
}

NUMBER_FIELDS = {"credit", "hours_total", "hours_week"}


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).split())


def convert(source, target):
    workbook = openpyxl.load_workbook(source, read_only=True)
    if "课程清单" in workbook.sheetnames:
        sheet = workbook["课程清单"]
    else:
        sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [clean(cell) for cell in next(rows)]
    index = {}
    for position, title in enumerate(header):
        if title in COLUMNS:
            index[COLUMNS[title]] = position
    missing = [field for field in ("id", "term", "course_name", "teacher_name") if field not in index]
    if missing:
        raise SystemExit(f"表头缺少必要列: {', '.join(missing)}")

    courses = []
    seen = set()
    for row in rows:
        record = {}
        for field, position in index.items():
            value = clean(row[position] if position < len(row) else None)
            if field in NUMBER_FIELDS and value:
                try:
                    number = float(value)
                    value = int(number) if number == int(number) else number
                except ValueError:
                    value = ""
            record[field] = value
        if not record["id"]:
            continue
        if record["id"] in seen:
            raise SystemExit(f"教学班号重复: {record['id']}")
        if not record["course_name"] or not record["teacher_name"]:
            raise SystemExit(f"教学班 {record['id']} 缺少课程名称或教师")
        seen.add(record["id"])
        courses.append(record)

    if not courses:
        raise SystemExit("没有解析到任何课程记录")
    terms = {course["term"] for course in courses}
    snapshot = {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "term": courses[0]["term"] if len(terms) == 1 else "",
        "courses": courses,
    }
    with open(target, "w", encoding="utf-8") as output:
        json.dump(snapshot, output, ensure_ascii=False, indent=1)
    print(f"converted {len(courses)} courses -> {target} (term={snapshot['term'] or 'mixed'})")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("用法: python3 xlsx_to_courses.py 输入.xlsx 输出.json")
    convert(sys.argv[1], sys.argv[2])
